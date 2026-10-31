"""Builds the specification set: SRS (IEEE 830), SAD (IEEE 1016) and Software Test Plan (IEEE 829).

    python tools/build_specification.py

Writes docs/reports/source/SRS_STP_SAD_Report.docx and docs/reports/1. SRS STP SAD Report.pdf.
Requirement links, measured values and test results are read from the code, traceability.py and the
latest test run in docs/reports/test-reports/raw, so the document matches the system.
The PDF and the page numbers in the table of contents need LibreOffice (soffice) and pdftotext.
"""
# pylint: disable=too-many-lines  # the module holds the text of a 35-page document

from __future__ import annotations

import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Dict, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)
sys.path.insert(0, os.path.join(BASE_DIR, "tools"))

# pylint: disable=wrong-import-position  # project modules need the paths above
import generate_test_reports as gtr  # noqa: E402
import report_diagrams as diagrams  # noqa: E402
from specification_content import REQUIREMENTS, SEQ_BATCH, SEQ_BENCH, SEQ_PREDICT, SEQ_TRAIN  # noqa: E402
import prediction_pipeline as pp  # noqa: E402
from traceability import TEST_CASES, requirement_index  # noqa: E402

DOCX_OUT = os.path.join(BASE_DIR, "docs", "reports", "source", "SRS_STP_SAD_Report.docx")
PDF_OUT = os.path.join(BASE_DIR, "docs", "reports", "1. SRS STP SAD Report.pdf")
DOC_ID = "HF-CCPS-SPEC"
VERSION = "2.0"
TITLE = "Customer Churn Prediction System"
FONT = "Calibri"
WIDTH_CM = 16.0  # text width on A4 with 2.5 cm margins

# ----- document helper ------------------------------------------------------------------------
class Spec:  # pylint: disable=too-many-public-methods
    """python-docx wrapper with numbered captions, standard tables and table-of-contents bookkeeping."""

    def __init__(self, chart_dir: str):
        self.doc = Document()
        self.chart_dir = chart_dir
        self.headings: List[Tuple[int, str]] = []
        self.figures = 0
        self.tables = 0
        self._styles()

    def _styles(self) -> None:
        styles = self.doc.styles
        for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet",
                     "List Number", "Caption"]:
            style = styles[name]
            style.font.name = FONT
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        styles["Normal"].font.size = Pt(10.5)
        styles["Normal"].paragraph_format.space_after = Pt(5)
        styles["Normal"].paragraph_format.line_spacing = 1.12
        for name, size, before in [("Heading 1", 15, 16), ("Heading 2", 12.5, 12), ("Heading 3", 11, 9)]:
            styles[name].font.size = Pt(size)
            styles[name].font.bold = True
            styles[name].paragraph_format.space_before = Pt(before)
            styles[name].paragraph_format.space_after = Pt(5)
            styles[name].paragraph_format.keep_with_next = True
        styles["Caption"].font.size = Pt(9)
        styles["Caption"].font.italic = False
        styles["Caption"].font.bold = True
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(2.5))
        section.top_margin, section.bottom_margin = Cm(2.3), Cm(2.2)
        section.different_first_page_header_footer = True
        self._header_footer(section)

    @staticmethod
    def _field(run, code: str) -> None:
        for tag, text in (("begin", None), (None, code), ("end", None)):
            if tag:
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), tag)
            else:
                element = OxmlElement("w:instrText")
                element.set(qn("xml:space"), "preserve")
                element.text = text
            run._r.append(element)  # pylint: disable=protected-access  # python-docx has no field API

    def _header_footer(self, section) -> None:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header.add_run(f"{TITLE}: SRS, SAD and STP   |   {DOC_ID}   |   Version {VERSION}")
        run.font.size = Pt(8.5)
        _bottom_border(header)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Team Hackforge, PES University   |   Page ")
        run.font.size = Pt(8.5)
        for code, suffix in (("PAGE", " of "), ("NUMPAGES", "")):
            field_run = footer.add_run()
            field_run.font.size = Pt(8.5)
            self._field(field_run, code)
            if suffix:
                footer.add_run(suffix).font.size = Pt(8.5)

    # content
    def heading(self, level: int, text: str, toc: bool = True) -> None:
        """Numbered heading; levels 1 and 2 go into the table of contents."""
        self.doc.add_heading(text, level)
        if toc and level <= 2:
            self.headings.append((level, text))

    def part(self, text: str) -> None:
        """Starts a part (A, B, C) on a new page."""
        self.page_break()
        self.heading(1, text)

    def para(self, text: str, bold: Optional[str] = None, italic: bool = False, align=None) -> None:
        """Paragraph, optionally starting with a bold lead-in."""
        paragraph = self.doc.add_paragraph()
        if bold:
            paragraph.add_run(bold).bold = True
        run = paragraph.add_run(text)
        run.italic = italic
        if align is not None:
            paragraph.alignment = align

    def bullets(self, items: Sequence[str], numbered: bool = False) -> None:
        """Bulleted or numbered list."""
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            if ": " in item and len(item.split(": ")[0]) < 45:
                lead, rest = item.split(": ", 1)
                paragraph.add_run(lead + ": ").bold = True
                paragraph.add_run(rest)
            else:
                paragraph.add_run(item)

    def table(self, headers: Sequence[str], rows: Sequence[Sequence[object]], widths: Sequence[float],
              caption: Optional[str] = None, font: float = 9.0) -> None:
        """Ruled table with a shaded header row and an optional numbered caption above it."""
        # pylint: disable=too-many-locals
        if caption:
            self.tables += 1
            caption_paragraph = self.doc.add_paragraph(f"Table {self.tables}. {caption}", style="Caption")
            caption_paragraph.paragraph_format.keep_with_next = True
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        scale = WIDTH_CM / sum(widths)
        for column, width in zip(table.columns, widths):
            column.width = Cm(width * scale)
        _repeat_header(table.rows[0])
        for cell, header, width in zip(table.rows[0].cells, headers, widths):
            cell.width = Cm(width * scale)
            cell.paragraphs[0].add_run(str(header)).bold = True
            _shade(cell, "E6E6E6")
        for row in rows:
            cells = table.add_row().cells
            for cell, value, width in zip(cells, row, widths):
                cell.width = Cm(width * scale)
                cell.paragraphs[0].add_run(str(value))
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.size = Pt(font)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def kv(self, rows: Sequence[Tuple[str, str]], widths: Tuple[float, float] = (4.2, 11.8),
           caption: Optional[str] = None) -> None:
        """Two-column key/value table without a header row."""
        if caption:
            self.tables += 1
            self.doc.add_paragraph(f"Table {self.tables}. {caption}", style="Caption")
        table = self.doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        for key, value in rows:
            cells = table.add_row().cells
            for cell, text, width in zip(cells, (key, value), widths):
                cell.width = Cm(width)
                run = cell.paragraphs[0].add_run(text)
                run.font.size = Pt(9)
                cell.paragraphs[0].paragraph_format.space_after = Pt(1)
            cells[0].paragraphs[0].runs[0].bold = True
            _shade(cells[0], "F2F2F2")
        _keep_together(table)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, path: str, caption: str, width: float = 15.5) -> None:
        """Centered image with a numbered caption below."""
        self.figures += 1
        self.doc.add_picture(path, width=Cm(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        caption_paragraph = self.doc.add_paragraph(f"Figure {self.figures}. {caption}", style="Caption")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def chart(self, name: str) -> str:
        """Path for a generated diagram."""
        return os.path.join(self.chart_dir, name)

    def requirements(self, ids: Sequence[str], caption: Optional[str] = None) -> None:
        """Requirement table: ID, statement, rank and verifying test cases."""
        index = requirement_index()
        rows = [[rid, REQUIREMENTS[rid][0], REQUIREMENTS[rid][1], ", ".join(index.get(rid, []))] for rid in ids]
        self.table(["ID", "Requirement", "Rank", "Verified by"], rows, [1.6, 10.6, 1.9, 1.9], caption)

    def page_break(self) -> None:
        """Hard page break."""
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self, entries: List[Tuple[int, str]], pages: Dict[int, int]) -> None:
        """Static table of contents with dot leaders; page numbers come from the rendered PDF."""
        self.heading(1, "Table of Contents", toc=False)
        for i, (level, text) in enumerate(entries):
            paragraph = self.doc.add_paragraph()
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT,
                                                              WD_TAB_LEADER.DOTS)
            paragraph.paragraph_format.left_indent = Cm(0 if level == 1 else 0.6)
            paragraph.paragraph_format.space_after = Pt(1 if level == 2 else 2)
            paragraph.paragraph_format.space_before = Pt(5 if level == 1 else 0)
            run = paragraph.add_run(f"{text}\t{pages.get(i, '')}")
            run.font.size = Pt(10 if level == 1 else 9.5)
            run.bold = level == 1


def _bottom_border(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()  # pylint: disable=protected-access
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    border.append(bottom)
    properties.append(border)


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()  # pylint: disable=protected-access
    shading = OxmlElement("w:shd")
    for key, value in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        shading.set(qn(key), value)
    properties.append(shading)


def _keep_together(table) -> None:
    """Rows never split across pages, and the whole table stays on one page with the heading above it."""
    for index, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()  # pylint: disable=protected-access
        properties.append(OxmlElement("w:cantSplit"))
        if index < len(table.rows) - 1:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()  # pylint: disable=protected-access
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


# ----- data -------------------------------------------------------------------------------------
def _test_plan_cases() -> Dict[str, Dict[str, str]]:
    """Steps, expected results and automated tests per test case from Test Plan v1.1, section 11.1."""
    cases = {}
    with open(os.path.join(BASE_DIR, "docs", "TEST_PLAN_v1.1.md"), encoding="utf-8") as handle:
        for line in handle:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if len(cells) == 7 and re.fullmatch(r"TC-\d\d", cells[0]):
                cases[cells[0]] = {"steps": cells[3], "expected": cells[4], "tests": cells[5].replace("`", "")}
    return cases


def _results() -> Dict[str, object]:  # pylint: disable=too-many-locals
    junit = gtr.read_junit()
    coverage = gtr.read_coverage()
    with open(os.path.join(gtr.RAW_DIR, "pylint_output.txt"), encoding="utf-8") as handle:
        match = re.search(r"rated at ([0-9.]+)/10", handle.read())
    pylint = float(match.group(1)) if match else 0.0
    cases = gtr._case_results(junit, pylint)  # pylint: disable=protected-access
    totals = coverage["totals"]
    measured = {k: v for r in cases.values() for k, v in r["measured"].items()}
    return {
        "junit": junit, "cases": cases, "pylint": pylint, "measured": measured,
        "tests": len(junit["tests"]), "passed": sum(t["outcome"] == "Passed" for t in junit["tests"]),
        "statement": 100 * totals["covered_lines"] / totals["num_statements"],
        "combined": 100 * (totals["covered_lines"] + totals["covered_branches"])
                    / (totals["num_statements"] + totals["num_branches"]),
        "branch": 100 * totals["covered_branches"] / totals["num_branches"],
        "date": junit["timestamp"][:10],
    }


def _paradigms() -> Dict[str, Dict[str, float]]:
    with open(os.path.join(BASE_DIR, "paradigm_metrics.json"), encoding="utf-8") as handle:
        return json.load(handle)


# ----- front matter -----------------------------------------------------------------------------
def front_matter(spec: Spec) -> None:
    """Cover page and document control."""
    doc = spec.doc
    for _ in range(6):
        doc.add_paragraph()
    for text, style in [(TITLE, "Title"), ("Software Requirements Specification (SRS)", "Subtitle"),
                        ("Software Architecture and Design Description (SAD)", "Subtitle"),
                        ("Software Test Plan (STP)", "Subtitle")]:
        paragraph = doc.add_paragraph(text, style=style)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spec.para("Prepared in accordance with IEEE Std 830-1998, IEEE Std 1016-2009 and IEEE Std 829-2008",
              italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4):
        doc.add_paragraph()
    spec.kv([("Document identifier", DOC_ID), ("Version", f"{VERSION} (released for evaluation)"),
             ("Date", dt.date.today().strftime("%d %B %Y")), ("Prepared by", "Team Hackforge"),
             ("Institution", "PES University, Software Engineering Mini Project"),
             ("Instructor", "Cvsn Reddy")], widths=(4.5, 11.5))
    spec.page_break()

    spec.heading(1, "Document Control", toc=False)
    spec.heading(2, "Document Information", toc=False)
    spec.kv([
        ("Title", f"{TITLE}: Software Requirements Specification, Software Architecture and Design Description, "
                  "and Software Test Plan"),
        ("Identifier", DOC_ID), ("Version", VERSION), ("Status", "Released for evaluation"),
        ("Classification", "Academic, public repository"),
        ("Repository", "github.com/Akshayshetty742/hackforge-customer-churn"),
        ("Source of this document", "tools/build_specification.py (regenerated from the code and the test run)"),
    ])
    spec.heading(2, "Project Team", toc=False)
    spec.table(["Name", "SRN", "Role in this document"], [
        ["Akshaya B", "PES1UG24AM026", "Author and reviewer"],
        ["Ashwin C H", "PES1UG24AM900", "Author and reviewer"],
        ["Amar Patil", "PES1UG24AM031", "Author and reviewer"],
    ], [5, 4, 7])
    spec.heading(2, "Revision History", toc=False)
    spec.table(["Version", "Date", "Description"], [
        ["1.0", "12 Oct 2026", "Initial SRS, SAD and Test Plan for the command-line training pipeline "
                               "(Deliverable 1)."],
        ["1.1", "14 Oct 2026", "Added the Streamlit web application, input validation and test cases TC-14 to "
                               "TC-21."],
        ["2.0", dt.date.today().strftime("%d %b %Y"),
         "Consolidated release after Sprints 1 to 3: batch scoring, risk explanation and what-if simulation, "
         "ML vs DML vs QML comparison, DVC and MLflow, 26 automated test cases with execution results. "
         "Restructured to the full IEEE 830, 1016 and 829 outlines."],
    ], [2.0, 2.8, 11.2])
    spec.heading(2, "Approvals", toc=False)
    spec.table(["Role", "Name", "Signature", "Date"], [
        ["Prepared by", "Team Hackforge", "", ""], ["Reviewed by", "", "", ""], ["Approved by", "Cvsn Reddy", "", ""],
    ], [3.5, 5, 4.5, 3])


# ----- Part A: SRS --------------------------------------------------------------------------------
def part_a(spec: Spec) -> None:  # pylint: disable=too-many-statements
    """Software Requirements Specification (IEEE 830-1998)."""
    spec.part("Part A: Software Requirements Specification")
    spec.heading(1, "1. Introduction")
    spec.heading(2, "1.1 Purpose")
    spec.para("This Software Requirements Specification (SRS) defines the functional, interface, performance, "
              "quality and security requirements of the Customer Churn Prediction System (CCPS). It is the "
              "agreed baseline for design (Part B), verification (Part C) and acceptance. The intended audience is "
              "the development team, the course instructor and reviewers, and anyone maintaining or extending "
              "the system.")
    spec.heading(2, "1.2 Scope")
    spec.para("CCPS predicts whether a telecom customer is likely to cancel the service (churn), explains the main "
              "risk factors and supports retention decisions. It comprises three subsystems:")
    spec.bullets([
        "Training pipeline: loads and prepares the Telco Customer Churn dataset, trains and tunes classical "
        "models, evaluates them on a held-out partition and persists the selected model with its encoders.",
        "Web application: single-customer prediction, batch scoring, model evaluation with an adjustable decision "
        "threshold, risk explanation and what-if retention simulation.",
        "Paradigm benchmark: trains classical machine learning (XGBoost), deep learning (PyTorch MLP) and quantum "
        "machine learning (PennyLane variational quantum classifier) models on an identical split, tracks every "
        "run in MLflow and serves all three predictions in the application.",
    ])
    spec.para("CCPS is an academic decision-support system. It does not contact customers, integrate with "
              "production CRM or billing systems, or guarantee that a flagged customer will churn. All processing "
              "is local to the user's machine.")
    spec.heading(2, "1.3 Definitions, Acronyms and Abbreviations")
    spec.table(["Term", "Definition"], [
        ["Churn", "A customer cancelling or not renewing the service; the target attribute Churn (Yes / No)."],
        ["CCPS", "Customer Churn Prediction System, the system specified in this document."],
        ["ML / DML / QML", "Classical machine learning / deep machine learning / quantum machine learning."],
        ["XGBoost", "Gradient-boosted decision tree library used for the production model."],
        ["MLP", "Multi-layer perceptron, the deep learning model."],
        ["VQC", "Variational quantum classifier, the quantum machine learning model."],
        ["Precision", "Of the customers predicted to churn, the fraction that actually churned."],
        ["Recall", "Of the customers that actually churned, the fraction predicted to churn."],
        ["ROC-AUC", "Area under the receiver operating characteristic curve; ranking quality across thresholds."],
        ["Decision threshold", "Churn probability at or above which a customer is classified as Churn (default "
                               "0.50)."],
        ["Risk tier", "Low (probability below 35%), Moderate (35% to 65%) or High (above 65%)."],
        ["Held-out partition", "Stratified 20% test split (random seed 42, 1,409 customers) never used for "
                               "training or tuning."],
        ["DVC", "Data Version Control: dataset versioning and reproducible pipeline stages."],
        ["MLflow", "Experiment tracking: parameters, metrics, artifacts and models of every run."],
        ["CI", "Continuous integration (GitHub Actions)."],
        ["Essential / Conditional", "IEEE 830 ranking: essential requirements must be met for acceptance; "
                                    "conditional requirements enhance the product."],
    ], [3.6, 12.4], caption="Definitions and abbreviations")
    spec.heading(2, "1.4 References")
    spec.bullets([
        "[1] IEEE Std 830-1998, IEEE Recommended Practice for Software Requirements Specifications.",
        "[2] IEEE Std 1016-2009, IEEE Standard for Information Technology: Software Design Descriptions.",
        "[3] IEEE Std 829-2008, IEEE Standard for Software and System Test Documentation.",
        "[4] IBM Telco Customer Churn sample dataset (WA_Fn-UseC_-Telco-Customer-Churn.csv).",
        "[5] Project repository: github.com/Akshayshetty742/hackforge-customer-churn (README, DEVELOPER_GUIDE, "
        "PROJECT_REPORT, docs/reports).",
        "[6] Library documentation: scikit-learn, XGBoost, PyTorch, PennyLane, Streamlit, MLflow, DVC.",
    ])
    spec.heading(2, "1.5 Overview")
    spec.para("Section 2 of this part describes the product context, functions, users, constraints and "
              "assumptions. Section 3 states the specific requirements: external interfaces, functional "
              "requirements organised by feature, performance, logical data, design constraints and software system "
              "attributes, including security. Each requirement has a unique identifier, an IEEE 830 rank and the "
              "test case that verifies it (Part C).")

    spec.heading(1, "2. Overall Description")
    spec.heading(2, "2.1 Product Perspective")
    spec.para("CCPS is a self-contained product that runs on a single workstation. It consumes the Telco customer "
              "dataset, produces model artifacts and experiment records, and serves predictions through a local "
              "web interface. The source code, continuous integration and backlog are hosted on GitHub.")
    spec.figure(diagrams.context_diagram(spec.chart("context.png")), "System context diagram")
    spec.heading(3, "2.1.1 System Interfaces")
    spec.para("CCPS has no runtime interface to external systems. It reads a CSV file, writes model and metric "
              "files, and stores experiment records in a local SQLite database used by MLflow.")
    spec.heading(3, "2.1.2 User Interfaces")
    spec.para("A browser-based interface (Streamlit) with four pages: Customer Churn Prediction, Batch Customer "
              "Scoring, Model Evaluation and About the Project. Command-line entry points are provided for training, "
              "benchmarking, testing and report generation (Part B, Section 3.6).")
    spec.heading(3, "2.1.3 Hardware Interfaces")
    spec.para("None beyond a standard workstation. No GPU or quantum hardware is required; the quantum model runs "
              "on a state-vector simulator.")
    spec.heading(3, "2.1.4 Software Interfaces")
    spec.table(["Software", "Purpose", "Minimum version"], [
        ["Python", "Runtime", "3.10 (3.12 used in CI)"],
        ["pandas, NumPy", "Tabular data handling", "2.0 / 1.24"],
        ["scikit-learn", "Preprocessing, cross-validation, metrics, estimator interface", "1.4"],
        ["XGBoost", "Production classifier", "2.0"],
        ["Streamlit, Altair", "Web application and charts", "1.36 / 5.0"],
        ["PyTorch", "Deep learning model", "2.2"],
        ["PennyLane", "Quantum circuit simulation", "0.38"],
        ["MLflow", "Experiment tracking", "3.0"],
        ["DVC", "Data and pipeline versioning", "3.0"],
    ], [3.5, 8.5, 4.0], caption="Software interfaces")
    spec.heading(3, "2.1.5 Communications Interfaces")
    spec.para("HTTP on the local loopback interface only: port 8501 for the web application and port 5001 for the "
              "MLflow user interface (port 5000 is reserved by macOS AirPlay). No data is sent to external hosts.")
    spec.heading(3, "2.1.6 Memory Constraints")
    spec.para("The dataset (7,043 records) and all models fit comfortably in the memory of a standard laptop. "
              "The serialized models total about 0.3 MB.")
    spec.heading(3, "2.1.7 Operations")
    spec.bullets(["Training mode: batch execution of the training pipeline or the paradigm benchmark.",
                  "Interactive mode: single-customer prediction, evaluation and simulation in the browser.",
                  "Batch mode: scoring of a customer file through the web application."])
    spec.heading(3, "2.1.8 Site Adaptation Requirements")
    spec.para("All training settings (data split, hyperparameter grid, network and circuit sizes) are defined in "
              "params.yaml. Ports can be changed with the standard Streamlit and MLflow command-line options.")

    spec.heading(2, "2.2 Product Functions")
    spec.bullets([
        "Load, validate and prepare the customer dataset.",
        "Train, tune, evaluate and persist the production churn model.",
        "Validate a customer profile and predict churn probability, class and risk tier.",
        "Compare classical, deep learning and quantum model predictions for the same customer.",
        "Evaluate the production model on the held-out partition at a chosen decision threshold.",
        "Score a batch of customers and export the results.",
        "Explain risk drivers and simulate retention actions.",
        "Benchmark the three paradigms and track every experiment.",
    ])
    spec.figure(diagrams.use_case_diagram(spec.chart("use_cases.png")), "Use case diagram")
    spec.table(["Use case", "Primary actor", "Description", "Requirements"], [
        ["UC-01 Load dataset", "Data scientist", "Read the dataset and confirm required columns.", "FR-01, FR-02"],
        ["UC-02 Prepare features", "Data scientist", "Separate the target and encode categorical attributes.",
         "FR-03, FR-04"],
        ["UC-03 Train model", "Data scientist", "Train and tune the classifier with class weighting.",
         "FR-05, FR-06, FR-10, FR-11"],
        ["UC-04 Evaluate model", "Data scientist", "Compute metrics and the confusion matrix on held-out data.",
         "FR-07, FR-08, FR-12"],
        ["UC-05 Save artifacts", "Data scientist", "Persist the model with its encoders.", "FR-09, FR-14"],
        ["UC-06 Handle bad input", "(included)", "Reject missing or invalid data with a descriptive error.",
         "FR-02, FR-16"],
        ["UC-07 Predict churn", "Retention analyst", "Enter or load a profile and obtain a prediction.",
         "FR-15, FR-16, FR-17"],
        ["UC-08 Compare paradigms", "Retention analyst", "View ML, DML and QML predictions side by side.", "FR-21"],
        ["UC-09 Evaluate and tune", "Retention analyst", "Re-evaluate the model at a chosen threshold.", "FR-18"],
        ["UC-10 Score a batch", "Retention analyst", "Score a customer file and download the results.", "FR-19"],
        ["UC-11 Explain and simulate", "Retention analyst", "View risk drivers and simulate retention actions.",
         "FR-20"],
        ["UC-12 Benchmark paradigms", "Data scientist", "Train and compare ML, DML and QML; track runs.",
         "FR-22, FR-23, FR-24"],
    ], [3.6, 2.8, 6.4, 3.2], caption="Use case summary")

    spec.heading(2, "2.3 User Characteristics")
    spec.table(["User class", "Characteristics", "Main use cases"], [
        ["Retention analyst", "Business user; no programming knowledge; needs clear labels, explanations and "
                              "plain-language results.", "UC-07 to UC-11"],
        ["Data scientist / ML engineer", "Runs training and benchmarks from the command line; understands "
                                         "classification metrics.", "UC-01 to UC-05, UC-12"],
        ["Reviewer / instructor", "Inspects documentation, tests and results.", "All (read-only)"],
    ], [3.8, 8.6, 3.6], caption="User classes")
    spec.heading(2, "2.4 Constraints")
    spec.bullets([
        "Data: model quality is limited by the size and content of the public dataset (7,043 customers, 26.5% "
        "churners).",
        "Imbalance: accuracy alone is misleading; recall and precision shall always be reported (NFR-04).",
        "Locality: the system runs on one workstation and shall not transmit customer data (COM-01).",
        "Quantum execution: the quantum model runs on a classical simulator.",
        "Compatibility: a persisted model shall be used only with the encoders saved in the same run (NFR-07).",
    ])
    spec.heading(2, "2.5 Assumptions and Dependencies")
    spec.bullets([
        "The dataset is available at the configured path with the expected columns and Yes / No target values.",
        "Python and the libraries in the requirements files are installed in the active environment.",
        "Persisted models are loaded with library versions compatible with those used to create them.",
        "The deep and quantum features depend on PyTorch and PennyLane; without them the application continues "
        "to serve the classical prediction and explains which package is missing.",
    ])
    spec.heading(2, "2.6 Apportioning of Requirements")
    spec.para("The following capabilities are deferred to future releases: containerised deployment and a REST "
              "prediction API, production monitoring and drift detection, scheduled retraining, and execution of "
              "the quantum model on quantum hardware.")

    spec.heading(1, "3. Specific Requirements")
    spec.para("Requirements use \"shall\" for mandatory behaviour. Rank follows IEEE 830: Essential requirements "
              "must be met for acceptance; Conditional requirements enhance the product. \"Verified by\" lists the "
              "Part C test cases.")
    spec.heading(2, "3.1 External Interface Requirements")
    spec.heading(3, "3.1.1 User Interfaces")
    spec.requirements(["UI-01", "UI-02"])
    spec.heading(3, "3.1.2 Hardware Interfaces")
    spec.requirements(["HW-01"])
    spec.heading(3, "3.1.3 Software Interfaces")
    spec.requirements(["SW-01", "SW-02", "SW-03"])
    spec.heading(3, "3.1.4 Communications Interfaces")
    spec.requirements(["COM-01"])

    spec.heading(2, "3.2 Functional Requirements")
    features = [
        ("3.2.1 Data Ingestion and Preparation", "The data scientist starts the training pipeline. The system "
         "reads the dataset, checks the required columns and target, removes the customer identifier, separates "
         "the target and encodes categorical attributes.", ["FR-01", "FR-02", "FR-03", "FR-04"]),
        ("3.2.2 Model Training and Tuning", "The system trains the configured model families on the training "
         "partition, searches the hyperparameter grid with stratified cross-validation and weights the minority "
         "class.", ["FR-05", "FR-06", "FR-10", "FR-11"]),
        ("3.2.3 Evaluation and Reporting", "After training, the system scores the held-out partition once and "
         "records labelled metrics, the confusion matrix and plots.", ["FR-07", "FR-08", "FR-12"]),
        ("3.2.4 Model Persistence", "The selected model is saved with its encoders and can be reloaded to score new "
         "customers.", ["FR-09", "FR-14"]),
        ("3.2.5 Customer Profile Entry and Validation", "The analyst fills in the form or loads a sample customer "
         "and submits it. The system validates every field before any prediction.", ["FR-15", "FR-16"]),
        ("3.2.6 Churn Prediction", "For a valid profile the system encodes the record in training column order and "
         "returns the prediction.", ["FR-17"]),
        ("3.2.7 Model Evaluation Interface", "The analyst requests a live evaluation and adjusts the decision "
         "threshold to explore the precision and recall trade-off.", ["FR-18"]),
        ("3.2.8 Batch Scoring", "The analyst uploads a customer file or selects the sample batch; the system scores "
         "all records and offers a download.", ["FR-19"]),
        ("3.2.9 Risk Explanation and Retention Simulation", "The system lists the factors driving a customer's "
         "risk; the analyst changes retention levers and sees the updated probability.", ["FR-20"]),
        ("3.2.10 Multi-Paradigm Prediction and Benchmarking", "The data scientist runs the benchmark; the analyst "
         "sees the three paradigm predictions for each customer.", ["FR-21", "FR-22", "FR-23", "FR-24"]),
        ("3.2.11 Documentation", "The project documentation supports installation, operation and evaluation.",
         ["FR-13"]),
    ]
    for title, stimulus, ids in features:
        spec.heading(3, title)
        spec.para(stimulus, bold="Description and stimulus/response: ")
        spec.requirements(ids)

    spec.heading(2, "3.3 Performance Requirements")
    spec.requirements(["NFR-01", "NFR-11"], caption="Performance requirements")
    spec.heading(2, "3.4 Logical Database Requirements")
    spec.para("CCPS has no database management system. Its logical data consists of the customer record, the model "
              "artifacts and the experiment records. The data dictionary below defines the customer record and the "
              "validation rules applied to it (FR-16).")
    rows = []
    for column in pp.EXPECTED_FEATURE_COLUMNS:
        if column in pp.VALID_CATEGORIES:
            rows.append([column, pp.FIELD_LABELS.get(column, column), "Categorical",
                         ", ".join(pp.VALID_CATEGORIES[column])])
        else:
            low, high = pp.NUMERIC_RANGES[column]
            kind = "Integer" if column in ("SeniorCitizen", "tenure") else "Decimal"
            rows.append([column, pp.FIELD_LABELS.get(column, column), kind, f"{low:g} to {high:g} (inclusive)"])
    spec.table(["Attribute", "Label", "Type", "Permitted values"], rows, [3.4, 3.4, 2.2, 7.0],
               caption="Customer record data dictionary", font=8.5)
    low, high = pp.BILLING_RATIO_RANGE
    spec.para("Consistency rules: total charges shall lie between "
              f"{low:g} and {high:g} times monthly charges multiplied by tenure (and shall not exceed monthly "
              "charges when tenure is 0); customers without phone service shall have 'No phone service' for "
              "multiple lines; customers without internet service shall have 'No internet service' for every "
              "internet add-on, and customers with internet service shall not.")
    spec.table(["Data object", "Content", "Retention"], [
        ["Production model", "Dictionary {model, features_names} (customer_churn_model_xgboost.pkl)",
         "Versioned with the code"],
        ["Encoders", "One label encoder per categorical attribute and the target (encoders.pkl)",
         "Saved with the model"],
        ["Paradigm models", "Dictionary with model, features, paradigm, test metrics, parameter count "
                            "(models/*.pkl)", "Regenerated by the benchmark"],
        ["Metrics", "metrics.json, classification_report.json, paradigm_metrics.json", "Tracked by DVC"],
        ["Experiment records", "MLflow SQLite store (mlflow.db) with artifacts", "Local only"],
    ], [3.2, 9.0, 3.8], caption="Persistent data objects")
    spec.heading(2, "3.5 Design Constraints")
    spec.bullets([
        "Standards compliance: documentation follows IEEE 830, 1016 and 829; Python code follows PEP 8 as checked by "
        "pylint (NFR-13).",
        "The production model shall remain XGBoost with class weighting; the test partition shall never be used for "
        "tuning.",
        "Every model shall expose the scikit-learn estimator interface (fit, predict_proba, predict).",
        "New data shall be encoded with the persisted encoders before scoring.",
    ])
    spec.heading(2, "3.6 Software System Attributes")
    attributes = [
        ("3.6.1 Reliability", ["NFR-02", "NFR-10"]),
        ("3.6.2 Reproducibility", ["NFR-03", "NFR-12"]),
        ("3.6.3 Maintainability", ["NFR-05", "NFR-13"]),
        ("3.6.4 Portability", ["NFR-06"]),
        ("3.6.5 Usability", ["NFR-08"]),
        ("3.6.6 Integrity and Model Quality Reporting", ["NFR-04", "NFR-07"]),
    ]
    for title, ids in attributes:
        spec.heading(3, title)
        spec.requirements(ids)
    spec.heading(3, "3.6.7 Security")
    spec.para("Security objectives define what must be protected; security requirements define how.")
    spec.requirements(["SO-01", "SO-02", "SO-03"], caption="Security objectives")
    spec.requirements(["NFR-09", "SR-01", "SR-02", "SR-03", "SR-04", "SR-05", "SR-06", "SR-07"],
                      caption="Security and privacy requirements")
    spec.heading(2, "3.7 Other Requirements")
    spec.para("None. Requirement totals are summarised below.")
    groups = [("Interface", ("UI", "HW", "SW", "COM")), ("Functional", ("FR",)), ("Non-functional", ("NFR",)),
              ("Security objectives", ("SO",)), ("Security requirements", ("SR",))]
    rows = [[name, sum(1 for rid in REQUIREMENTS if rid.split("-")[0] in prefixes)] for name, prefixes in groups]
    rows.append(["Total", len(REQUIREMENTS)])
    spec.table(["Group", "Requirements"], rows, [8, 8], caption="Requirement summary")


# ----- Part B: SAD --------------------------------------------------------------------------------
def part_b(spec: Spec) -> None:  # pylint: disable=too-many-statements
    """Software Architecture and Design Description (IEEE 1016-2009 viewpoints)."""
    spec.part("Part B: Software Architecture and Design Description")
    spec.heading(1, "1. Introduction")
    spec.heading(2, "1.1 Purpose")
    spec.para("This part describes the architecture and detailed design of CCPS using the viewpoints of IEEE "
              "1016-2009. It shows how each requirement in Part A is realised and provides the basis for "
              "implementation, maintenance and verification.")
    spec.heading(2, "1.2 Scope")
    spec.para("The design covers the training pipeline, the prediction service, the web application, the three "
              "model paradigms, the benchmark and the MLOps tooling (DVC, MLflow, continuous integration).")
    spec.heading(2, "1.3 Design Overview")
    spec.para("CCPS uses a layered architecture. A single service module (prediction_pipeline.py) holds all "
              "business rules and is shared by the web application, the training scripts and the tests. Training is "
              "organised as a pipeline whose stages are reproducible through DVC and recorded in MLflow. The deep "
              "and quantum models are adapted to the scikit-learn estimator interface, so one evaluation, logging "
              "and serving path serves all paradigms.")

    spec.heading(1, "2. Stakeholders and Design Concerns")
    spec.table(["Stakeholder", "Design concerns", "Addressed in"], [
        ["Retention analyst", "Clear inputs and results, fast response, trustworthy predictions",
         "3.5, 3.6, NFR-11"],
        ["Data scientist", "Fair model comparison, reproducible experiments, configurable training",
         "3.4, 3.5, 3.7"],
        ["Maintainer", "Separation of concerns, testability, code quality", "3.2, 3.3, 6"],
        ["Reviewer / instructor", "Traceability from requirements to design and tests", "7, Part C"],
        ["Data owner", "Privacy, local processing, artifact integrity", "4"],
    ], [3.4, 8.2, 4.4], caption="Stakeholders and concerns")

    spec.heading(1, "3. Design Views")
    spec.heading(2, "3.1 Context Viewpoint")
    spec.para("The system context is shown in Part A, Figure 1. CCPS interacts with two user classes and four "
              "external stores: the dataset, model artifacts, the experiment store and the GitHub repository with "
              "its CI service.")
    spec.heading(2, "3.2 Composition Viewpoint")
    spec.figure(diagrams.component_diagram(spec.chart("components.png")), "Component diagram")
    spec.table(["Component", "Responsibility", "Provided interface"], [
        ["app.py", "Four-page web application; layout, state and charts", "Browser UI"],
        ["customer_presets.py", "Four realistic sample customers", "get_preset_values()"],
        ["prediction_pipeline.py", "Validation, preprocessing, prediction, risk tiers, evaluation, batch scoring, "
                                   "explanation, paradigm model loading", "Python API (3.6)"],
        ["classical_models.py", "Logistic Regression, Random Forest and XGBoost trainers with grid search",
         "train_*()"],
        ["dml_model.py", "TorchTabularClassifier base and the MLP", "scikit-learn estimator"],
        ["qml_model.py", "Variational quantum classifier", "scikit-learn estimator"],
        ["training_common.py", "Configuration, stratified split, metrics, plots, MLflow logging",
         "prepare_data(), evaluate_and_log_run()"],
        ["01_churn_prediction_xgboost.py", "Production training pipeline (DVC stage train)", "CLI"],
        ["02_paradigm_benchmark.py", "ML vs DML vs QML benchmark (DVC stage benchmark_paradigms)", "CLI"],
    ], [4.4, 7.8, 3.8], caption="Component responsibilities")
    spec.heading(3, "3.2.1 Architectural Patterns")
    spec.table(["Pattern", "Application", "Rationale"], [
        ["Layered architecture", "Presentation, service, model, training and storage layers",
         "Models can change without affecting the user interface."],
        ["Pipes and filters", "Training and benchmark stages; dvc.yaml",
         "Each stage has one responsibility and defined inputs and outputs."],
        ["Adapter", "TorchTabularClassifier wraps PyTorch and PennyLane models",
         "All paradigms share the scikit-learn interface."],
        ["Template method", "_build_network() overridden by the MLP and the VQC",
         "Training, early stopping and persistence are written once."],
        ["Registry / strategy", "TRAINERS and build_plan() select models from params.yaml",
         "Model families change without code changes."],
        ["Cache", "Streamlit resource cache for models", "Models load once per session (NFR-11)."],
    ], [3.4, 6.2, 6.4], caption="Architectural patterns")

    spec.heading(2, "3.3 Logical Viewpoint")
    spec.figure(diagrams.class_diagram(spec.chart("classes.png")), "Class diagram of the model layer", width=14.5)
    spec.para("TorchTabularClassifier implements preprocessing, class-weighted binary cross-entropy training with "
              "Adam, early stopping that restores the epoch with the lowest validation loss, prediction and "
              "pickling (weights are stored as a state dictionary and the network is rebuilt on load). Subclasses "
              "only define their network.")

    spec.heading(2, "3.4 Information Viewpoint")
    spec.para("The customer record and persistent data objects are defined in Part A, Section 3.4. Data flows from "
              "the raw CSV through cleaning (customer ID removed, blank total charges set to 0), label encoding and "
              "a stratified 80/20 split (seed 42). The deep and quantum models additionally apply one-hot encoding "
              "and standard scaling (45 inputs); the quantum model reduces these to 6 principal components scaled to "
              "rotation angles in [0, pi].")

    spec.heading(2, "3.5 Interaction Viewpoint")
    spec.figure(diagrams.sequence_diagram(SEQ_TRAIN, spec.chart("seq_train.png")),
                "Sequence diagram: training, evaluation and persistence")
    spec.figure(diagrams.sequence_diagram(SEQ_PREDICT, spec.chart("seq_predict.png")),
                "Sequence diagram: single-customer prediction with three paradigms")
    spec.figure(diagrams.sequence_diagram(SEQ_BATCH, spec.chart("seq_batch.png")),
                "Sequence diagram: batch scoring", width=13.5)
    spec.figure(diagrams.sequence_diagram(SEQ_BENCH, spec.chart("seq_bench.png")),
                "Sequence diagram: paradigm benchmark")

    spec.heading(2, "3.6 Interface Viewpoint")
    spec.para("By design CCPS exposes no network API (COM-01). Its programming interface is the prediction_pipeline "
              "module, used by the application, the training scripts and the tests.")
    spec.table(["Function", "Input", "Output", "Errors"], [
        ["load_artifacts(model_path, encoder_path)", "File paths (defaults: production files)",
         "(model, feature_names, encoders)", "FileNotFoundError, ValueError"],
        ["validate_customer_input(record)", "Dictionary of 19 fields", "(is_valid, errors)", "None (errors returned)"],
        ["preprocess_input(record, encoders, features)", "Record or table", "Encoded table in training order",
         "ValueError"],
        ["predict_churn(model, row, threshold=0.5)", "Encoded row", "(class, probability, probabilities, tier)",
         "None"],
        ["get_risk_tier(probability)", "Probability", "Low, Moderate or High", "None"],
        ["evaluate_model_on_dataset(threshold=0.5)", "Optional paths and threshold",
         "Metrics, confusion matrix, report", "FileNotFoundError, ValueError"],
        ["compare_with_benchmark(results)", "Live results", "(all_match, comparison table)", "None"],
        ["score_customer_batch(table, model, encoders, features, threshold)", "Customer table",
         "Table with probability, class, tier", "ValueError"],
        ["explain_customer_risk(record, probability)", "Record", "Risk-raising and risk-lowering factors", "None"],
        ["load_paradigm_model(path, features)", "Model file path", "Validated model payload",
         "FileNotFoundError, ImportError, ValueError"],
        ["predict_all_paradigms(row, models, threshold)", "Encoded row, models", "One result per paradigm", "None"],
    ], [5.4, 3.4, 4.4, 2.8], caption="Prediction service API", font=8.3)
    spec.table(["Command", "Purpose"], [
        ["streamlit run app.py", "Start the web application (port 8501)"],
        ["python 01_churn_prediction_xgboost.py  (dvc repro train)", "Train the classical models"],
        ["python 02_paradigm_benchmark.py  (dvc repro benchmark_paradigms)", "Run the ML vs DML vs QML benchmark"],
        ["mlflow ui --backend-store-uri sqlite:///mlflow.db --port 5001", "Browse experiments"],
        ["python -m pytest", "Run the automated tests"],
        ["python tools/generate_test_reports.py", "Run the tests and generate the test reports"],
    ], [9.0, 7.0], caption="Command-line interface")

    spec.heading(2, "3.7 Algorithm Viewpoint")
    spec.table(["Paradigm", "Model", "Input processing", "Structure", "Training"], [
        ["ML", "XGBoost", "Label-encoded attributes", "Gradient-boosted trees; depth, learning rate and number of "
                                                      "trees from a grid search", "scale_pos_weight 2.77; 3-fold "
                                                                                  "stratified CV"],
        ["DML", "Multi-layer perceptron", "One-hot and standardised (45 inputs)", "64 and 32 units; Linear, "
                                                                                 "BatchNorm, ReLU, Dropout 0.3",
         "Adam (lr 0.001), class-weighted BCE, early stopping (patience 15)"],
        ["QML", "Variational quantum classifier", "PCA to 6 features scaled to [0, pi]", "6 qubits; angle "
                                                                                         "embedding re-uploaded "
                                                                                         "before each of 3 entangling "
                                                                                         "layers; Pauli-Z readout; "
                                                                                         "linear output (61 "
                                                                                         "parameters)",
         "Adam (lr 0.02), class-weighted BCE, early stopping (patience 8)"],
    ], [1.6, 2.6, 3.4, 4.6, 3.8], caption="Model design", font=8.5)
    spec.para("Champion selection: the training pipeline trains every configured family and exports the model with "
              "the highest value of model.champion_metric (default ROC-AUC). Imbalance correction: the positive "
              "class weight equals the ratio of non-churners to churners in the training partition (about 2.77).")

    spec.heading(2, "3.8 Deployment Viewpoint")
    spec.figure(diagrams.deployment_diagram(spec.chart("deployment.png")), "Deployment diagram")
    spec.para("The application, MLflow user interface and all data run on the user's workstation. GitHub hosts the "
              "repository; GitHub Actions runs the core tests and the quality job (pylint, full test suite with "
              "coverage, report artifacts) for every pull request and every push to main.")

    spec.heading(1, "4. Security Architecture")
    spec.para("Trust boundaries: the user's workstation (trusted), input files supplied by users (untrusted data) "
              "and the public repository (published content).")
    spec.table(["Threat", "Control", "Requirements"], [
        ["Unexpected or malicious values in input files", "Schema and value validation; files parsed as data only",
         "SR-05, SO-03"],
        ["Personal data used by the model", "Identifier removed before modelling; input columns fixed by the "
                                            "feature list", "SR-01, NFR-09"],
        ["Tampered or mismatched model file", "Models loaded only from models/; payload format and feature list "
                                              "checked before use", "SR-07, NFR-07"],
        ["Model and encoders from different runs", "Saved together; integrity verified by tests", "SR-03, SO-02"],
        ["Data leaving the machine", "No network clients in the code; local MLflow store", "COM-01, SR-04, SR-06"],
        ["Secrets committed to the repository", ".gitignore rules and an automated tracked-file check", "SR-02"],
        ["Production model overwritten by tests", "Tests redirect all outputs to temporary folders", "SO-02"],
    ], [4.6, 8.0, 3.4], caption="Threats and controls")

    spec.heading(1, "5. Error Handling Design")
    spec.table(["Condition", "System behaviour", "Component"], [
        ["Dataset missing or unreadable", "Stops with a message naming the file; no outputs written",
         "prediction_pipeline.load_dataset"],
        ["Required column or target missing", "Stops with a message listing the columns", "load_dataset"],
        ["Invalid form input", "Lists every invalid field; no prediction shown", "validate_customer_input"],
        ["Unknown category or missing column in a batch", "Error naming the problem; batch not scored",
         "preprocess_input"],
        ["params.yaml missing or invalid", "Documented defaults used; warning printed", "training_common.load_config"],
        ["Unknown model type or paradigm in configuration", "Stops listing valid options", "select_trainers, "
                                                                                          "build_plan"],
        ["PyTorch or PennyLane not installed", "Classical prediction served; card names the missing package",
         "app.py"],
        ["Paradigm model missing or incompatible", "That card shows the reason; other results unaffected",
         "load_paradigm_model"],
    ], [4.8, 6.8, 4.4], caption="Error handling")

    spec.heading(1, "6. Design Rationale")
    spec.table(["Decision", "Alternatives considered", "Rationale"], [
        ["Class weighting for imbalance", "SMOTE oversampling, downsampling", "SMOTE overfitted; downsampling "
                                                                              "discarded data; weighting keeps every "
                                                                              "record."],
        ["XGBoost as production model", "Random Forest, Logistic Regression, MLP", "Best ROC-AUC on the held-out "
                                                                                   "set with built-in "
                                                                                   "regularisation."],
        ["Streamlit for the interface", "Flask with a custom front end", "Fast delivery of an interactive, "
                                                                          "testable (AppTest) interface."],
        ["scikit-learn interface for all models", "Separate code paths per paradigm", "One evaluation, logging and "
                                                                                       "serving path; fair "
                                                                                       "comparison."],
        ["DVC and MLflow", "Manual experiment records", "Reproducible stages and automatic, complete run records."],
        ["Single-threaded PyTorch inference on macOS", "Default threading", "Prevents an OpenMP conflict between "
                                                                            "PyTorch and XGBoost."],
        ["Local-only processing", "Hosted service", "Protects customer data; meets COM-01."],
    ], [4.2, 4.4, 7.4], caption="Design decisions")

    spec.heading(1, "7. Requirements Traceability")
    spec.table(["Requirements", "Design element"], [
        ["FR-01 to FR-04, SW-01", "training_common.prepare_data, prediction_pipeline.load_dataset"],
        ["FR-05, FR-06, FR-10, FR-11", "classical_models (grid search, scale_pos_weight)"],
        ["FR-07, FR-08, FR-12, UI-01, NFR-04", "training_common.evaluate_and_log_run, metrics.json"],
        ["FR-09, FR-14, SW-02, NFR-07", "01_churn_prediction_xgboost.export_champion, load_artifacts"],
        ["FR-15, FR-16, UI-02, NFR-08", "app.py form, customer_presets, validate_customer_input"],
        ["FR-17, FR-18, FR-19, FR-20", "predict_churn, evaluate_model_on_dataset, score_customer_batch, "
                                       "explain_customer_risk"],
        ["FR-21 to FR-24, SR-07", "dml_model, qml_model, 02_paradigm_benchmark, predict_all_paradigms"],
        ["NFR-01, NFR-11, HW-01", "Grid size in params.yaml; Streamlit resource cache"],
        ["NFR-03, NFR-12, SW-03, SR-06", "Fixed seeds, dvc.yaml, params.yaml, MLflow logging"],
        ["NFR-05, NFR-06, NFR-13", "Module structure, requirements files, pyproject.toml, CI workflow"],
        ["COM-01, SR-01 to SR-05, SO-01 to SO-03, NFR-09, NFR-02, NFR-10", "Security architecture (Section 4) and "
                                                                           "error handling (Section 5)"],
    ], [6.0, 10.0], caption="Requirements to design traceability")


# ----- Part C: STP --------------------------------------------------------------------------------
def part_c(spec: Spec, results: Dict[str, object]) -> None:  # pylint: disable=too-many-statements,too-many-locals
    """Software Test Plan (IEEE 829-2008)."""
    measured = results["measured"]
    spec.part("Part C: Software Test Plan")
    spec.heading(1, "1. Introduction")
    spec.heading(2, "1.1 Test Plan Identifier")
    spec.kv([("Identifier", "STP-CCPS-2.0"), ("Items under test", "CCPS release 2.0 (main branch)"),
             ("Related documents", "Part A (SRS), Part B (SAD), Unit Test Report, System Test Report, coverage "
                                   "reports")])
    spec.heading(2, "1.2 Purpose and Scope")
    spec.para("This plan defines the scope, approach, resources and schedule of testing for CCPS, the test cases "
              "that verify every requirement in Part A, and the results of their execution.")
    spec.heading(2, "1.3 References")
    spec.para("IEEE Std 829-2008; Part A; Part B; test reports in docs/reports/test-reports.")

    spec.heading(1, "2. Test Items")
    spec.table(["Test item", "Files"], [
        ["Prediction service", "prediction_pipeline.py, customer_presets.py"],
        ["Web application", "app.py"],
        ["Training pipeline", "01_churn_prediction_xgboost.py, training_common.py, classical_models.py"],
        ["Deep and quantum models", "dml_model.py, qml_model.py"],
        ["Paradigm benchmark", "02_paradigm_benchmark.py"],
        ["Configuration", "params.yaml, dvc.yaml, requirements*.txt, .gitignore"],
        ["Documentation", "README.md, DEVELOPER_GUIDE.md, docs/"],
    ], [4.5, 11.5], caption="Test items")

    spec.heading(1, "3. Features to be Tested")
    rows = [[case_id, case.title, ", ".join(case.requirements)] for case_id, case in TEST_CASES.items()]
    spec.table(["Test case", "Feature", "Requirements"], rows, [2.2, 7.0, 6.8], caption="Features to be tested",
               font=8.5)

    spec.heading(1, "4. Features Not to be Tested")
    spec.table(["Feature", "Reason"], [
        ["Whether a flagged customer will actually churn", "Future behaviour cannot be verified in testing."],
        ["Retention campaigns, CRM integration", "Outside the system scope (Part A, 1.2)."],
        ["Internal correctness of third-party libraries", "Vendor responsibility; correct use is tested."],
        ["Quantum hardware execution", "The system targets a simulator."],
        ["Formal security certification", "Outside the academic scope."],
    ], [7.0, 9.0], caption="Features not to be tested")

    spec.heading(1, "5. Approach")
    spec.heading(2, "5.1 Security Validation")
    spec.table(["Check", "Method", "Pass criterion", "Test case"], [
        ["Data minimisation", "Inspect model inputs after preparation", "No identifier (ID, name, e-mail, phone, "
                                                                         "address) is an input", "TC-11"],
        ["Repository hygiene", "Check .gitignore and tracked files", "Secrets and environments ignored; no secret "
                                                                     "files tracked", "TC-11"],
        ["Local processing", "Scan source for network clients; inspect MLflow store", "No network client; local "
                                                                                       "store only",
         "TC-11, TC-23"],
        ["Input handling", "Missing file, missing target, invalid values, unknown categories", "Controlled error; no "
                                                                                               "result reported",
         "TC-02, TC-15"],
        ["Artifact integrity", "Reload model with encoders; load invalid paradigm models", "Pairs load and predict; "
                                                                                           "invalid files rejected",
         "TC-09, TC-20, TC-25"],
    ], [3.0, 5.0, 5.4, 2.6], caption="Security validation")
    spec.heading(2, "5.2 Test Levels")
    spec.table(["Level", "Scope", "Tools"], [
        ["Unit", "Validation rules, metrics, risk tiers, model classes", "pytest"],
        ["Integration", "Preparation with training; persistence and reload; preprocessing with prediction", "pytest"],
        ["System", "Full training run, full benchmark, web application run headless", "pytest, Streamlit AppTest"],
        ["Non-functional", "Runtime, latency, repeatability, portability, privacy, documentation, code quality",
         "pytest, pylint, coverage.py"],
    ], [3.0, 9.4, 3.6], caption="Test levels")
    spec.heading(2, "5.3 Test Design Techniques")
    spec.bullets([
        "Equivalence partitioning: valid and invalid classes for each categorical and numeric field.",
        "Boundary value analysis: numeric limits (for example tenure -1, 0, 120, 121) and threshold boundaries.",
        "Decision tables: phone and internet service consistency rules.",
        "End-to-end scenarios: training and benchmark pipelines executed in isolated temporary folders.",
        "Analytical oracles: a one-qubit circuit compared with its closed-form result; metrics on known inputs.",
    ])
    spec.heading(2, "5.4 Automation and Continuous Integration")
    spec.para("All test cases are automated with pytest and tagged with their test case identifier. GitHub "
              "Actions runs the core tests with the application requirements only, and a quality job with pylint, "
              "the full suite and statement and branch coverage, for every pull request and every push to main.")
    spec.heading(2, "5.5 Model Evaluation Procedure")
    spec.bullets([
        "Use the stratified 20% held-out partition (seed 42, 1,409 customers) for every model.",
        "Confirm churn is the positive class and that the partition is never used for tuning.",
        "Compute accuracy, precision, recall, F1-score, ROC-AUC and the confusion matrix.",
        "Compare live results of the production model with the recorded benchmark (tolerance 0.005).",
    ], numbered=True)
    spec.heading(2, "5.6 Test Data")
    spec.para("The project dataset for end-to-end tests; generated temporary CSV files for invalid-input tests; "
              "the four sample customers for prediction tests; small data slices for fast model unit tests.")

    spec.heading(1, "6. Item Pass/Fail Criteria")
    spec.bullets([
        "A test case passes when all of its automated tests pass.",
        "The release passes when every test case passes, combined statement and branch coverage of the application "
        "modules is at least 85%, and the pylint score is at least 8.5/10.",
        "Severity: Critical (privacy breach, target leakage, test data used for tuning); Major (a required function, "
        "metric or artifact fails); Minor (documentation or presentation defect).",
    ])
    spec.heading(1, "7. Suspension Criteria and Resumption Requirements")
    spec.para("Testing is suspended if the dataset or production model is missing, the environment cannot be "
              "installed, or a privacy or data-leakage problem is found. After correction the full suite is "
              "executed again.")
    spec.heading(1, "8. Test Deliverables")
    spec.table(["Deliverable", "Location"], [
        ["Software Test Plan (this part)", "docs/reports/1. SRS STP SAD Report.pdf"],
        ["Unit Test Report", "docs/reports/5. Unit Test Report.pdf"],
        ["System Test Report", "docs/reports/4. System Test Report.pdf"],
        ["Code and Branch Coverage Reports", "docs/reports/6 and 7"],
        ["Raw evidence (JUnit XML, coverage XML/JSON, pytest and pylint output)", "docs/reports/test-reports/raw"],
    ], [8.0, 8.0], caption="Test deliverables")
    spec.heading(1, "9. Testing Tasks")
    spec.bullets(["Maintain test cases and their requirement links (traceability.py).",
                  "Automate each test case and tag it with its identifier.",
                  "Execute the suite locally and in CI; generate the reports.",
                  "Record defects, correct them and re-execute the suite."], numbered=True)
    spec.heading(1, "10. Environmental Needs")
    spec.table(["Item", "Value"], [list(row) for row in gtr.environment()], [6, 10], caption="Test environment",
               font=8.5)
    spec.heading(1, "11. Responsibilities")
    spec.table(["Role", "Responsibility"], [
        ["Developer", "Write and automate test cases for each backlog item; keep the suite passing."],
        ["Reviewer", "Review pull requests; confirm each new feature has a linked test case."],
        ["Instructor", "Review the documentation, reports and demonstration."],
    ], [4, 12])
    spec.heading(1, "12. Staffing and Training Needs")
    spec.para("Three team members. Required skills: Python, pytest, scikit-learn and Streamlit testing; training "
              "material is provided in DEVELOPER_GUIDE.md.")
    spec.heading(1, "13. Schedule")
    spec.table(["Activity", "Date"], [
        ["Test plan v1.0 (TC-01 to TC-13)", "12 October 2026"],
        ["Web application test cases (TC-14 to TC-21)", "14 October 2026"],
        ["Sprint 1 and Sprint 2 automated tests and CI", "October 2026"],
        ["Test plan update and automation of all 26 test cases", "30 October 2026"],
        ["Testing week and report regeneration (Deliverable Part 3)", "2 to 6 November 2026"],
    ], [10, 6], caption="Test schedule")
    spec.heading(1, "14. Risks and Contingencies")
    spec.table(["Risk", "Contingency"], [
        ["Library update changes model results", "TC-17 and TC-25 detect it; regenerate model and encoders together."],
        ["Optional packages unavailable", "Dependent tests skip; CI quality job installs all packages."],
        ["Tests overwrite the production model", "All outputs redirected to temporary folders."],
        ["Documentation drifts from the code", "Reports and this document are generated from the code and test run."],
    ], [6, 10], caption="Test risks")
    spec.heading(1, "15. Approvals")
    spec.table(["Role", "Name", "Signature", "Date"], [["Test lead", "", "", ""], ["Reviewer", "", "", ""],
                                                        ["Instructor", "Cvsn Reddy", "", ""]], [3.5, 5, 4.5, 3])

    spec.heading(1, "16. Test Case Specifications")
    plan = _test_plan_cases()
    cases = results["cases"]
    for case_id, case in TEST_CASES.items():
        spec.heading(3, f"{case_id} {case.title}", toc=False)
        detail = plan.get(case_id, {})
        result = cases[case_id]
        outcome = (f"{sum(t == 'Passed' for t in result['tests'])} of {len(result['tests'])} automated tests passed"
                   if result["tests"] else f"pylint score {results['pylint']:.2f}/10")
        spec.kv([("Requirements", ", ".join(case.requirements)), ("Level", case.level),
                 ("Steps", detail.get("steps", "")), ("Expected result", detail.get("expected", "")),
                 ("Automated by", detail.get("tests", "")), ("Result", f"{result['verdict']} ({outcome})")],
                widths=(3.4, 12.6))

    spec.heading(1, "17. Requirements Traceability Matrix")
    index = requirement_index()
    order = ["UI", "HW", "SW", "COM", "FR", "NFR", "SO", "SR"]
    ids = sorted(index, key=lambda r: (order.index(r.split("-")[0]), int(r.split("-")[1])))
    verdict = {cid: cases[cid]["verdict"] for cid in TEST_CASES}
    rows = [[rid, ", ".join(index[rid]), "Verified" if all(verdict[c] == "Pass" for c in index[rid]) else "Open"]
            for rid in ids]
    half = (len(rows) + 1) // 2
    paired = [left + (right if right else ["", "", ""]) for left, right in
              zip(rows[:half], rows[half:] + [None] * (half - len(rows[half:])))]
    spec.table(["Req.", "Test cases", "Status", "Req.", "Test cases", "Status"], paired,
               [1.8, 3.4, 1.8, 1.8, 3.4, 1.8], caption="Requirements traceability matrix", font=8.3)

    spec.heading(1, "18. Test Execution Summary")
    passed_cases = sum(r["verdict"] == "Pass" for r in cases.values())
    spec.table(["Measure", "Target", "Result"], [
        ["Test cases passed", "All", f"{passed_cases} of {len(cases)}"],
        ["Automated tests passed", "All", f"{results['passed']} of {results['tests']}"],
        ["Requirements verified", "All", f"{len(ids)} of {len(REQUIREMENTS)}"],
        ["Combined statement and branch coverage", ">= 85%", f"{results['combined']:.1f}%"],
        ["Statement coverage (application modules)", "Reported", f"{results['statement']:.1f}%"],
        ["Branch coverage (application modules)", "Reported", f"{results['branch']:.1f}%"],
        ["pylint score", ">= 8.5", f"{results['pylint']:.2f}"],
        ["Training time (NFR-01)", "< 300 s", f"{measured.get('training_seconds', 'n/a')} s"],
        ["Prediction latency, median (NFR-11)", "< 200 ms", f"{measured.get('median_prediction_ms', 'n/a')} ms"],
    ], [7.0, 4.0, 5.0], caption=f"Execution summary (test run of {results['date']})")
    benchmark = pp.SAVED_BENCHMARK_METRICS
    spec.para(f"Production model on the held-out partition: accuracy {benchmark['test_accuracy']:.2%}, "
              f"precision {benchmark['test_precision']:.2%}, recall {benchmark['test_recall']:.2%}, F1 "
              f"{benchmark['test_f1']:.3f}, ROC-AUC {benchmark['test_roc_auc']:.3f}.")
    spec.table(["Paradigm", "Model", "Accuracy", "Recall", "ROC-AUC", "Parameters"], [
        [m["paradigm"], name, f"{m['test_accuracy']:.2%}", f"{m['churn_recall']:.2%}", f"{m['roc_auc']:.3f}",
         f"{int(m['trainable_parameters']):,}"] for name, m in _paradigms().items()],
        [2.0, 4.0, 2.4, 2.4, 2.4, 2.8], caption="Paradigm benchmark results")
    spec.table(["ID", "Defect", "Found by", "Resolution", "Status"], [
        ["DEF-01", "app.py imported altair, which no requirements file declared", "TC-10",
         "Declared altair in requirements.txt", "Closed"],
        ["DEF-02", "Process could hang on macOS when PyTorch ran multi-threaded after XGBoost", "Benchmark run",
         "PyTorch inference runs single-threaded", "Closed"],
        ["DEF-03", "XGBoost model failed to load when PyTorch was imported first (macOS OpenMP)", "TC-21 run",
         "XGBoost imported before PyTorch in dml_model.py", "Closed"],
        ["DEF-04", "Training exported the last trained model instead of the best one", "Code inspection",
         "Champion selected by model.champion_metric", "Closed"],
        ["DEF-05", "Early stopping could restore an epoch other than the best one", "TC-21",
         "Exact minimum validation loss tracked", "Closed"],
    ], [1.6, 5.6, 2.4, 4.6, 1.8], caption="Defects found and resolved", font=8.5)
    spec.para(f"Conclusion: all {len(cases)} test cases passed and all {len(REQUIREMENTS)} requirements are "
              "verified. CCPS release 2.0 meets the acceptance criteria of this plan.", bold="")


# ----- build --------------------------------------------------------------------------------------
def build(entries: Optional[List[Tuple[int, str]]], pages: Dict[int, int], chart_dir: str,
          results: Dict[str, object]) -> Spec:
    """Assembles the whole document."""
    spec = Spec(chart_dir)
    front_matter(spec)
    spec.page_break()
    spec.toc(entries or [], pages)
    part_a(spec)
    part_b(spec)
    part_c(spec, results)
    return spec


def _pdf(docx_path: str, out_dir: str) -> str:
    subprocess.run([_tool("soffice"), "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
                   check=True, capture_output=True)
    return os.path.join(out_dir, os.path.basename(docx_path).replace(".docx", ".pdf"))


def _tool(name: str) -> str:
    path = shutil.which(name) or shutil.which(name, path="/opt/homebrew/bin:/usr/local/bin")
    if not path:
        raise SystemExit(f"{name} is required (LibreOffice / poppler).")
    return path


def _page_numbers(pdf_path: str, entries: List[Tuple[int, str]]) -> Dict[int, int]:
    """Page of each table-of-contents entry, searched in order after the contents pages."""
    text = subprocess.run([_tool("pdftotext"), "-layout", pdf_path, "-"], check=True, capture_output=True,
                          text=True).stdout
    pages = [" ".join(page.split()) for page in text.split("\f")]
    start = next(i for i, page in enumerate(pages) if "Table of Contents" in page) + 1
    found: Dict[int, int] = {}
    cursor = start
    for i, (_, heading) in enumerate(entries):
        target = " ".join(heading.split())
        for page_index in range(cursor, len(pages)):
            if target in pages[page_index] and "Table of Contents" not in pages[page_index]:
                found[i] = page_index + 1
                cursor = page_index
                break
    return found


def main() -> None:
    """Two passes: render once to find heading pages, then write the final document with page numbers."""
    work = tempfile.mkdtemp()
    results = _results()
    entries = build(None, {}, work, results).headings
    first = build(entries, {}, work, results)
    draft = os.path.join(work, "draft.docx")
    first.doc.save(draft)
    pages = _page_numbers(_pdf(draft, work), entries)
    final = build(entries, pages, work, results)
    os.makedirs(os.path.dirname(DOCX_OUT), exist_ok=True)
    final.doc.save(DOCX_OUT)
    produced = _pdf(DOCX_OUT, work)
    shutil.move(produced, PDF_OUT)
    print(f"Wrote {os.path.relpath(DOCX_OUT, BASE_DIR)} and {os.path.relpath(PDF_OUT, BASE_DIR)} "
          f"({len(entries)} contents entries, {len(pages)} located)")


if __name__ == "__main__":
    main()
